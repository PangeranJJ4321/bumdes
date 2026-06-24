import csv
import docx
import urllib.request
import io
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    """Set cell padding (in dxa)"""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('w:top', top), ('w:bottom', bottom), ('w:left', left), ('w:right', right)]:
        node = OxmlElement(m)
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_cell_borders(cell):
    """Set standard black borders for the cell"""
    tcPr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>\n'
        f'  <w:top w:val="single" w:sz="4" w:space="0" w:color="000000"/>\n'
        # Left and right borders for inner cells
        f'  <w:left w:val="single" w:sz="4" w:space="0" w:color="000000"/>\n'
        f'  <w:bottom w:val="single" w:sz="4" w:space="0" w:color="000000"/>\n'
        f'  <w:right w:val="single" w:sz="4" w:space="0" w:color="000000"/>\n'
        f'</w:tcBorders>'
    )
    tcPr.append(borders)

def main():
    doc = docx.Document()
    
    # Set standard margins for A4 (1 inch = 72 pt = 1440 dxa)
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)
        section.page_width = Inches(8.27)  # A4 width
        section.page_height = Inches(11.69) # A4 height
        
    # Configure default style (Times New Roman, 11pt)
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Times New Roman'
    font.size = Pt(11)
    
    # Add a title
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_title = p_title.add_run("LOGBOOK HARIAN KKN TEMATIK")
    run_title.bold = True
    run_title.font.size = Pt(14)
    
    # Create Table
    table = doc.add_table(rows=1, cols=4)
    table.autofit = False
    
    # Table headers
    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = "NO."
    hdr_cells[1].text = "HARI/TGL"
    hdr_cells[2].text = "JAM"
    hdr_cells[3].text = "KEGIATAN PER HARI"
    
    # Style headers
    for i, cell in enumerate(hdr_cells):
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in cell.paragraphs[0].runs:
            run.bold = True
            run.font.size = Pt(11)
        set_cell_borders(cell)
        set_cell_margins(cell, top=120, bottom=120, left=150, right=150)
        
    # Column widths
    col_widths = [Inches(0.6), Inches(1.5), Inches(1.2), Inches(3.7)]
    for i, col in enumerate(table.columns):
        col.width = col_widths[i]
        
    # Set header row width
    for i, cell in enumerate(hdr_cells):
        cell.width = col_widths[i]

    # Read data from CSV
    with open('data_logbook.csv', mode='r', encoding='utf-8') as f:
        reader = csv.reader(f)
        header = next(reader)  # Skip header row
        
        for row in reader:
            if not row:
                continue
            
            # Columns: No, Judul, Tanggal, Rincian, Dokumentasi
            no_val = row[0].strip()
            judul_val = row[1].strip()
            tanggal_val = row[2].strip()
            rincian_val = row[3].strip() if len(row) > 3 else ""
            dok_val = row[4].strip() if len(row) > 4 else ""
            
            # Parse Tanggal
            hari_tgl = tanggal_val
            jam = ""
            if "Jam:" in tanggal_val:
                parts = tanggal_val.split("Jam:")
                hari_tgl = parts[0].strip()
                after_jam = parts[1].strip()
                if "JKEM:" in after_jam:
                    jam = after_jam.split("JKEM:")[0].strip()
                else:
                    jam = after_jam
            
            # Add table row
            row_cells = table.add_row().cells
            
            # Set column widths for this row
            for i, cell in enumerate(row_cells):
                cell.width = col_widths[i]
                set_cell_borders(cell)
                set_cell_margins(cell, top=100, bottom=100, left=150, right=150)
                
            # Column 0: NO.
            p_no = row_cells[0].paragraphs[0]
            p_no.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_no.add_run(no_val)
            
            # Column 1: HARI/TGL
            p_tgl = row_cells[1].paragraphs[0]
            p_tgl.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_tgl.add_run(hari_tgl)
            
            # Column 2: JAM
            p_jam = row_cells[2].paragraphs[0]
            p_jam.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_jam.add_run(jam)
            
            # Column 3: KEGIATAN PER HARI
            cell_kegiatan = row_cells[3]
            
            # Judul Kegiatan Section
            p_jud_lbl = cell_kegiatan.paragraphs[0]
            p_jud_lbl.add_run("Judul Kegiatan:").bold = True
            
            p_jud_bullet = cell_kegiatan.add_paragraph(style='List Bullet')
            p_jud_bullet.paragraph_format.space_after = Pt(6)
            p_jud_bullet.add_run(judul_val)
            
            # Rincian Kegiatan Section
            p_rinc_lbl = cell_kegiatan.add_paragraph()
            p_rinc_lbl.add_run("Rincian Kegiatan:").bold = True
            
            p_rinc_bullet = cell_kegiatan.add_paragraph(style='List Bullet')
            p_rinc_bullet.paragraph_format.space_after = Pt(6)
            p_rinc_bullet.add_run(rincian_val)
            
            # Dokumen Pendukung Section
            p_dok_lbl = cell_kegiatan.add_paragraph()
            p_dok_lbl.add_run("Dokumen Pendukung:").bold = True
            
            if dok_val:
                print(f"Downloading image for row {no_val}...")
                image_stream = None
                try:
                    req = urllib.request.Request(
                        dok_val, 
                        headers={'User-Agent': 'Mozilla/5.0'}
                    )
                    with urllib.request.urlopen(req, timeout=10) as response:
                        image_bytes = response.read()
                    image_stream = io.BytesIO(image_bytes)
                except Exception as e:
                    print(f"Failed to download image for row {no_val}: {e}")
                
                p_dok_val = cell_kegiatan.add_paragraph()
                p_dok_val.paragraph_format.space_after = Pt(12)
                if image_stream:
                    try:
                        p_img = cell_kegiatan.add_paragraph()
                        p_img.paragraph_format.space_after = Pt(12)
                        run_img = p_img.add_run()
                        run_img.add_picture(image_stream, width=Inches(3.0))
                    except Exception as img_err:
                        print(f"Failed to embed image for row {no_val}: {img_err}")
                        p_dok_val.add_run("Link gambar - ")
                        p_dok_val.add_run(dok_val).font.color.rgb = docx.shared.RGBColor(0, 0, 255)
                else:
                    p_dok_val.add_run("Link gambar - ")
                    p_dok_val.add_run(dok_val).font.color.rgb = docx.shared.RGBColor(0, 0, 255)
            else:
                p_dok_val = cell_kegiatan.add_paragraph()
                p_dok_val.paragraph_format.space_after = Pt(12)
                p_dok_val.add_run("Tidak ada dokumentasi")
                
    # Save the document
    try:
        doc.save('logbook_kkn.docx')
        print("Logbook successfully generated as logbook_kkn.docx")
    except PermissionError:
        doc.save('logbook_kkn_baru.docx')
        print("logbook_kkn.docx is currently locked. Saved as logbook_kkn_baru.docx instead.")

if __name__ == '__main__':
    main()
